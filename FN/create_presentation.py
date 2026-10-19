# ==============================================================================
# FINAL TERM: PROFESSIONAL 16:9 POWERPOINT PRESENTATION GENERATOR
# Real-Time Speech Command Recognition for Mobile Robot Edge Control (BC-ResNet)
# Covers 100% of Syllabus Points in 1pro.png & 2pro.png
# ==============================================================================
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES_DIR = os.path.join(BASE_DIR, "FN", "results")
PPTX_PATH = os.path.join(RES_DIR, "BC_RESNET_FINAL_REPORT.pptx")
GRADCAM_DIR = os.path.join(RES_DIR, "gradcam")

# Palette constants
C_NAVY_BG   = RGBColor(15, 23, 42)     # #0F172A (Deep Slate Navy)
C_CARD_BG   = RGBColor(248, 250, 252) # #F8FAFC (Clean Off-White Card)
C_DARK_CARD = RGBColor(30, 41, 59)    # #1E293B (Darker Card)
C_TEXT_DARK = RGBColor(15, 23, 42)    # #0F172A
C_TEXT_MUTED= RGBColor(100, 116, 139) # #64748B
C_TEXT_WHITE= RGBColor(255, 255, 255)
C_ACCENT_BL = RGBColor(37, 99, 235)   # #2563EB (Royal Blue)
C_EMERALD   = RGBColor(16, 185, 129)  # #10B981 (Emerald Green)
C_AMBER     = RGBColor(245, 158, 11)  # #F59E0B (Amber Gold)
C_BORDER    = RGBColor(226, 232, 240) # #E2E8F0 (Subtle Border)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, tag_text, title_text, subtitle_text):
    # Tag Badge
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(5.0), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = C_ACCENT_BL

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.55))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(20)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_DARK

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.22), Inches(11.7), Inches(0.4))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = subtitle_text
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = C_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
    else:
        card.line.fill.background()
    return card

# ==============================================================================
# SLIDE 1: TITLE SLIDE (Dark Navy Premium Design)
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_BG
bg1.line.fill.background()

# Title text
t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.333), Inches(2.2))
tf = t_box.text_frame
tf.word_wrap = True

p0 = tf.paragraphs[0]
p0.text = "ADVANCED DEEP LEARNING • FINAL TERM PROJECT DEFENSE"
p0.font.size = Pt(12)
p0.font.bold = True
p0.font.color.rgb = C_AMBER
p0.space_after = Pt(14)

p1 = tf.add_paragraph()
p1.text = "SOTA Broadcasted Residual Network (BC-ResNet)\nfor Mobile Robot Edge Keyword Spotting"
p1.font.size = Pt(28)
p1.font.bold = True
p1.font.color.rgb = C_TEXT_WHITE
p1.space_after = Pt(12)

p2 = tf.add_paragraph()
p2.text = "End-to-End Pipeline: Acoustic Preprocessing, SpecAugment, Frequency-Separable Convolutions,\nBroadcast Attention, Label Smoothing, and WebAssembly / BLE Edge Deployment"
p2.font.size = Pt(13)
p2.font.color.rgb = RGBColor(203, 213, 225)

# Student info
info_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(11.333), Inches(0.8))
tf_info = info_box.text_frame
p_info = tf_info.paragraphs[0]
p_info.text = "Student: Nguyen Hai Duong   |   Student ID: 23011418   |   Dataset: Data_FN (2,800 Balanced WAVs)"
p_info.font.size = Pt(13)
p_info.font.bold = True
p_info.font.color.rgb = C_TEXT_WHITE

# 4 Metric Highlights Badges
metrics = [
    ("TEST ACCURACY", "98.21%", "275 / 280 Unseen Samples Correct", C_EMERALD),
    ("PARAMETERS", "157.4 K", "48.8% Smaller than Midterm M5", C_ACCENT_BL),
    ("CPU INFERENCE", "4.78 ms", "Real-Time on Raspberry Pi / Web", C_AMBER),
    ("STANDALONE ONNX", "617 KB", "Opset 17 WebAssembly Deploy", RGBColor(168, 85, 247))
]

for idx, (m_tag, m_val, m_desc, m_col) in enumerate(metrics):
    x = Inches(1.0 + idx * 2.9)
    card = add_card(s1, x, Inches(5.3), Inches(2.7), Inches(1.5), bg_color=C_DARK_CARD, border_color=m_col)
    tb = s1.shapes.add_textbox(x, Inches(5.4), Inches(2.7), Inches(1.3))
    tf_m = tb.text_frame
    tf_m.word_wrap = True
    
    pm1 = tf_m.paragraphs[0]
    pm1.text = m_tag
    pm1.font.size = Pt(9)
    pm1.font.bold = True
    pm1.font.color.rgb = m_col
    
    pm2 = tf_m.add_paragraph()
    pm2.text = m_val
    pm2.font.size = Pt(22)
    pm2.font.bold = True
    pm2.font.color.rgb = C_TEXT_WHITE
    
    pm3 = tf_m.add_paragraph()
    pm3.text = m_desc
    pm3.font.size = Pt(9)
    pm3.font.color.rgb = RGBColor(148, 163, 184)

# ==============================================================================
# SLIDE 2: GLOBAL ARCHITECTURAL BLUEPRINT
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Stage 0: Global Architecture Blueprint",
           "End-to-End System Pipeline Architecture",
           "From 16 kHz Audio Capture to Acoustic Preprocessing, BC-ResNet Inference, and BLE Robot Kinematics")

# Left Column (Highlights)
add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.1))
t_col = s2.shapes.add_textbox(Inches(0.95), Inches(1.95), Inches(3.3), Inches(4.8))
tf_col = t_col.text_frame
tf_col.word_wrap = True

sections = [
    ("1. Acoustic Ingestion", "16,000 Hz, 16-bit Mono PCM WAV audio.\nSymmetric 2.0s Center-Pad eliminates phoneme clipping."),
    ("2. Log-Mel Features", "128 mel frequency bins × 63 time frames.\nExtracts rich vocal formant structures."),
    ("3. Data Augmentation", "Google Brain SpecAugment (p=0.8, F=16, T=12).\nApplied only during training (zero leakage)."),
    ("4. SOTA BC-ResNet", "4 stages of Broadcasted Residual Blocks.\nCuts parameters down to 157K with 4.78 ms latency."),
    ("5. Edge & Robotics", "Exported standalone ONNX model (617 KB).\nDeploys via WebAssembly and BLE telemetry.")
]
for i, (head, body) in enumerate(sections):
    p_h = tf_col.add_paragraph() if i > 0 else tf_col.paragraphs[0]
    p_h.text = f"• {head}"
    p_h.font.size = Pt(11)
    p_h.font.bold = True
    p_h.font.color.rgb = C_ACCENT_BL
    
    p_b = tf_col.add_paragraph()
    p_b.text = body
    p_b.font.size = Pt(9.5)
    p_b.font.color.rgb = C_TEXT_MUTED
    p_b.space_after = Pt(8)

# Right: Diagram
add_card(s2, Inches(4.7), Inches(1.8), Inches(7.8), Inches(5.1))
img_path = os.path.join(RES_DIR, "arch_BC_ResNet_Architecture.png")
if os.path.exists(img_path):
    s2.shapes.add_picture(img_path, Inches(4.85), Inches(2.2), width=Inches(7.5))

# ==============================================================================
# SLIDE 3: ACOUSTIC PREPROCESSING & SPECAUGMENT (Ch. 2.5 & 4.1)
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Stage 1 & 2: Signal Conditioning & Augmentation",
           "Standardized 2.0s Audio Waveforms & SpecAugment",
           "Ensuring Uniform Temporal Alignments and Frequency-Time Invariant Representations")

add_card(s3, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
img_path = os.path.join(RES_DIR, "audio_pipeline_specaugment.png")
if os.path.exists(img_path):
    s3.shapes.add_picture(img_path, Inches(0.95), Inches(1.95), width=Inches(11.4))

# Bullet highlights under image
h_box = s3.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.3))
tf_h = h_box.text_frame
tf_h.word_wrap = True
bullet_items = [
    "Symmetric Center-Padding (Eq. 1): Standardizes all variable-duration speech recordings into exactly 32,000 samples (2.0s at 16 kHz), preventing prefix or suffix syllable loss.",
    "Log Mel-Spectrogram (Eq. 2-3): Transforms raw waveforms into 128 Mel frequency bins × 63 time frames using STFT with Hann window (N_FFT=1024, Hop=512).",
    "SpecAugment (p=0.8): Applies stochastic Frequency Masking (F=16) and Time Masking (T=12) on training batches; strictly disabled on validation and testing."
]
for i, txt in enumerate(bullet_items):
    p = tf_h.paragraphs[0] if i == 0 else tf_h.add_paragraph()
    p.text = f"✔ {txt}"
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(4)

# ==============================================================================
# SLIDE 4: SOTA BC-RESNET ARCHITECTURE BLUEPRINT (Ch. 2.1 & 2.2)
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Stage 3: Deep Neural Network Architecture",
           "BC-ResNet: Qualcomm Broadcasted Residual Network",
           "Ultra-Compact 157K Parameter Model Designed for Micro-Edge & Real-Time Embedded Systems")

add_card(s4, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
img_path = os.path.join(RES_DIR, "arch_BC_ResNet_Architecture.png")
if os.path.exists(img_path):
    s4.shapes.add_picture(img_path, Inches(0.95), Inches(2.0), width=Inches(11.4))

h_box4 = s4.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.2))
tf_h4 = h_box4.text_frame
tf_h4.word_wrap = True
points = [
    "Stem Layer: Conv2d (3×3, stride 1) with 32 filters + BatchNorm2d + ReLU, producing an initial acoustic feature map of shape (32, 128, 63).",
    "4 Residual Stages: Sequentially processes features with channel expansion (32 → 64 → 128 → 128) and frequency/temporal decimation (strides of 2 in Stages 2 & 3).",
    "Head: Global Average Pooling (GAP) reduces spatial dimensions to a 128-d vector, followed by Dropout (p=0.15) and a Linear classifier (128 → 7 classes)."
]
for i, txt in enumerate(points):
    p = tf_h4.paragraphs[0] if i == 0 else tf_h4.add_paragraph()
    p.text = f"• {txt}"
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(3)

# ==============================================================================
# SLIDE 5: BC-RESBLOCK & BROADCAST ATTENTION (Ch. 2.1 & 2.2)
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Stage 3 Deep Dive: Broadcast Residual Attention",
           "Internal Mechanism of the BC-ResBlock",
           "Frequency-wise Separable Convolutions Coupled with Broadcast Frequency Attention (BFA)")

add_card(s5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
img_path = os.path.join(RES_DIR, "arch_BC_ResBlock_Mechanism.png")
if os.path.exists(img_path):
    s5.shapes.add_picture(img_path, Inches(0.95), Inches(2.0), width=Inches(11.4))

h_box5 = s5.shapes.add_textbox(Inches(1.0), Inches(5.4), Inches(11.3), Inches(1.3))
tf_h5 = h_box5.text_frame
tf_h5.word_wrap = True
points5 = [
    "Frequency-wise Separable Convolution: Uses a 3×1 1D Depthwise Conv along frequency bins (groups=C) followed by a 1×1 Pointwise Conv, drastically cutting FLOPs.",
    "Broadcast Frequency Attention (BFA): Average pools across frequency bins (dim=2) to produce z in R^(C' x 1 x T), applies 1×1 Conv + BN + Sigmoid to compute attention s_t.",
    "Broadcast Modulation: The temporal attention weights s_t in [0, 1] are broadcast-multiplied along the frequency axis F, dynamically recalibrating time-dependent phoneme bursts."
]
for i, txt in enumerate(points5):
    p = tf_h5.paragraphs[0] if i == 0 else tf_h5.add_paragraph()
    p.text = f"✔ {txt}"
    p.font.size = Pt(10)
    p.font.color.rgb = C_TEXT_DARK
    p.space_after = Pt(4)

# ==============================================================================
# SLIDE 6: OPTIMIZATION & LOSS DESIGN (Ch. 2.3 & 2.4)
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Stage 4: Optimization & Regularization Strategy",
           "Label Smoothing Cross-Entropy & Cosine Annealing",
           "Preventing Overconfidence & Ensuring Smooth Convergence in Low-Resource Edge Regimes")

add_card(s6, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
img_path = os.path.join(RES_DIR, "optimization_loss_lr_schedule.png")
if os.path.exists(img_path):
    s6.shapes.add_picture(img_path, Inches(1.1), Inches(2.0), width=Inches(11.1))

# ==============================================================================
# SLIDE 7: BENCHMARK EVOLUTION & PARETO FRONTIER (Ch. 2.2 & 2.5)
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Stage 4: Model Benchmarking & Pareto Trade-Off",
           "Architectural Evolution from Midterm to Final Term",
           "Empirical Demonstration: BC-ResNet Dominates the Pareto Optimal Frontier")

add_card(s7, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1))
img_path = os.path.join(RES_DIR, "benchmark_pareto_evolution.png")
if os.path.exists(img_path):
    s7.shapes.add_picture(img_path, Inches(1.0), Inches(1.95), width=Inches(11.3))

# ==============================================================================
# SLIDE 8: CLASSIFICATION DIAGNOSTICS & EVALUATION (Ch. 4.1 & 4.2)
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "Stage 4: Diagnostic Performance on Test Set",
           "Precision, Recall, F1-Score Breakdown & ROC-AUC Curves",
           "Rigorous Evaluation on 280 Completely Unseen Multi-Speaker Recordings (Data_FN)")

# Subplot 1: Metrics Breakdown
add_card(s8, Inches(0.8), Inches(1.8), Inches(7.5), Inches(5.1))
img_m = os.path.join(RES_DIR, "metrics_classification_breakdown.png")
if os.path.exists(img_m):
    s8.shapes.add_picture(img_m, Inches(0.9), Inches(1.95), width=Inches(7.3))

# Subplot 2: Confusion Matrix
add_card(s8, Inches(8.5), Inches(1.8), Inches(4.0), Inches(5.1))
img_cm = os.path.join(RES_DIR, "confusion_matrix_bc_resnet.png")
if os.path.exists(img_cm):
    s8.shapes.add_picture(img_cm, Inches(8.6), Inches(2.0), width=Inches(3.8))

# ==============================================================================
# SLIDE 9: EXPLAINABLE AI (XAI) VIA GRAD-CAM (Ch. 3.1, 3.2, 3.3)
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "Stage 4: Explainable AI (XAI) via Grad-CAM",
           "Acoustic Phoneme Localization & Decision Interpretability",
           "Validating that the Network Relies on Authentic Vocal Formants Rather than Ambient Background Noise")

# 4 Grad-CAM tiles: KhoiDong, Tien, TamDung, KhongXacDinh
gradcam_samples = [
    ("KhoiDong", "gradcam_KhoiDong.png"),
    ("Tien", "gradcam_Tien.png"),
    ("TamDung", "gradcam_TamDung.png"),
    ("KhongXacDinh", "gradcam_KhongXacDinh.png")
]

for idx, (cname, fname) in enumerate(gradcam_samples):
    row = idx // 2
    col = idx % 2
    x = Inches(0.8 + col * 5.95)
    y = Inches(1.8 + row * 2.55)
    
    add_card(s9, x, y, Inches(5.75), Inches(2.4))
    g_path = os.path.join(GRADCAM_DIR, fname)
    if os.path.exists(g_path):
        s9.shapes.add_picture(g_path, x + Inches(0.1), y + Inches(0.1), width=Inches(5.55))

# ==============================================================================
# SLIDE 10: LATENT EMBEDDING & DIGITAL TWIN DEPLOYMENT (Ch. 5 & 2.6)
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "Stage 5: Latent Space Representation & Edge Deployment",
           "Unsupervised Manifold Clustering & Live Robot Telemetry",
           "Translating SOTA Deep Learning into a Full-Stack Edge & WebAssembly Digital Twin")

# Left: Latent clustering (t-SNE & PCA)
add_card(s10, Inches(0.8), Inches(1.8), Inches(7.5), Inches(5.1))
img_lat = os.path.join(RES_DIR, "latent_space_tsne_pca.png")
if os.path.exists(img_lat):
    s10.shapes.add_picture(img_lat, Inches(0.9), Inches(1.95), width=Inches(7.3))

# Right: Deployment Summary Card
add_card(s10, Inches(8.5), Inches(1.8), Inches(4.0), Inches(5.1), bg_color=C_NAVY_BG, border_color=C_ACCENT_BL)
tb_d = s10.shapes.add_textbox(Inches(8.7), Inches(2.0), Inches(3.6), Inches(4.7))
tf_d = tb_d.text_frame
tf_d.word_wrap = True

dp1 = tf_d.paragraphs[0]
dp1.text = "EDGE & WEB DEPLOYMENT"
dp1.font.size = Pt(11)
dp1.font.bold = True
dp1.font.color.rgb = C_AMBER
dp1.space_after = Pt(10)

deploy_bullets = [
    ("ONNX Model Optimization", "Exported standalone ONNX model (Opset 17) with dynamic batching. File size: 617 KB."),
    ("WebAssembly In-Browser", "Zero-server, clientside inference via onnxruntime-web running directly on Chrome/Safari at 4.78 ms."),
    ("ESP32 BLE Robot Telemetry", "Commands mapped directly into Bluetooth Low Energy motor state machine (FWD, REV, LEFT, RIGHT, STOP)."),
    ("Live Interactive Dashboard", "Hosted globally on GitHub Pages with Push-to-Talk and live audio visualizers.")
]

for title, desc in deploy_bullets:
    p_t = tf_d.add_paragraph()
    p_t.text = f"✔ {title}"
    p_t.font.size = Pt(10)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_WHITE
    
    p_d = tf_d.add_paragraph()
    p_d.text = desc
    p_d.font.size = Pt(9)
    p_d.font.color.rgb = RGBColor(148, 163, 184)
    p_d.space_after = Pt(8)

# Save presentation
prs.save(PPTX_PATH)
print(f"[SUCCESS] PowerPoint presentation successfully generated: {PPTX_PATH}")

# Also copy to root Report directory if it exists
rep_dir = os.path.join(BASE_DIR, "Report")
if os.path.exists(rep_dir):
    import shutil
    shutil.copy2(PPTX_PATH, os.path.join(rep_dir, "BC_RESNET_FINAL_REPORT.pptx"))
    print(f"[SUCCESS] Copied to Report: {os.path.join(rep_dir, 'BC_RESNET_FINAL_REPORT.pptx')}")
